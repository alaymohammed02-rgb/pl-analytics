"""Writes the documentary's script as a Word file (right-to-left Arabic) and as Markdown, with scene times,
the narration (plain and vowelled), what is on screen, and the picture credits.

  python3 make_script.py [OUT_DIR]      needs build/timeline.json and build/voice/timing.json (render.py / tts.py)
"""
import json, os, sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor
from arabic import plain

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, 'build')

SCREEN = {
    's01': 'رغيف خبز وحبة دواء وبطارية تظهر مع ذكرها، ثم العنوان «الكيمياء» وقارورة تتغير ألوانها وجزيء ماء يدور.',
    's02': 'ثلاث صور أرشيفية بالأبيض والأسود: بيير وماري كوري في مختبرهما، وجدول مندلييف الأصلي (1869)، ومؤتمر سولفاي 1927.',
    's03': 'بطاقات: تركيب المادة / خصائصها / تفاعلاتها، ثم سلسلة الاسم، ثم بطاقة جابر بن حيان مع جهاز تقطير متحرك والتقطير والتبلور والتسامي.',
    's04': 'شجرة: «الكيمياء» في الوسط وخمسة فروع تمتد مع ذكر كل فرع.',
    's05': 'جزيء الأسبرين يدور، ثم بلورة كلوريد الصوديوم، ثم شريط DNA، مع بطاقات الأمثلة.',
    's06': 'أربع بطاقات بالتتابع: الصابون، والصدأ، وقشرة الخبز، والبطارية.',
    's07': 'أربع بطاقات علماء بصورهم: الرازي (رسم من مخطوطة لاتينية)، ولافوازييه (نقش)، ومندلييف وماري كوري (صورتان فوتوغرافيتان).',
    's08': 'ميدالية ذهبية وأربعة أسطر: منذ 1901، فانت هوف، أحمد زويل 1999، عمر ياغي 2025.',
    's09': '«خرافة أم حقيقة؟»، ثم الرصاص والذهب وعدد البروتونات، ثم قارورة الإكسير والزئبق، ثم العناصر القديمة الأربعة والجدول الدوري بـ118 عنصرًا.',
    's10': '«شكرًا على المتابعة» مع اسم المعد وفصله وأسطر الاعتمادات.',
}


def mm(x):
    return f'{int(x // 60)}:{int(x % 60):02d}'


def rtl_par(p, align=WD_ALIGN_PARAGRAPH.RIGHT):
    p._p.get_or_add_pPr().append(OxmlElement('w:bidi'))
    p.alignment = align
    return p


def run(p, text, size=12, bold=False, color=None, font='Arial'):
    r = p.add_run(text)
    r.font.size = Pt(size); r.bold = bold
    r.font.name = font
    rpr = r._r.get_or_add_rPr()
    rpr.append(OxmlElement('w:rtl'))
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.append(rf)
    rf.set(qn('w:cs'), font)
    if bold:
        rpr.append(OxmlElement('w:bCs'))
    if color:
        r.font.color.rgb = RGBColor(*color)
    return r


def main():
    script = json.load(open(os.path.join(HERE, 'script.json'), encoding='utf-8'))
    T = json.load(open(os.path.join(HERE, 'build', 'timeline.json')))
    timing = {x['id']: x for x in json.load(open(os.path.join(HERE, 'build', 'voice', 'timing.json'), encoding='utf-8'))}
    name = os.environ.get('DOC_NAME', 'زكريا أحمد العي'); cls = os.environ.get('DOC_CLASS', 'الفصل 302')
    arch = os.path.join(HERE, 'archive', 'archive.json')
    credits = json.load(open(arch, encoding='utf-8')) if os.path.exists(arch) else []
    used = {'archive_curie_1.jpg', 'archive_mendeleev_1.png', 'archive_solvay_1.jpg', 'portrait_razi_1.JPG',
            'portrait_lavoisier_1.jpg', 'portrait_mendeleev_2.jpg', 'portrait_curie_1.jpg'}
    credits = [c for c in credits if c['file'] in used]

    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8); sec.top_margin = sec.bottom_margin = Cm(1.8)
    # default (complex-script) font
    st = d.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:cs'), 'Arial')

    run(rtl_par(d.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER), 'فيلم وثائقي: الكيمياء', 24, True, (0x0B, 0x5C, 0x6B))
    run(rtl_par(d.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER),
        f'المدة {mm(T["total"])}  •  إعداد: {name}  •  {cls}', 12, False, (0x55, 0x55, 0x55))
    run(rtl_par(d.add_paragraph()), 'الصوت: تعليق عربي مولَّد بالذكاء الاصطناعي. الترجمة المكتوبة تظهر في أسفل الفيديو، وهي نفسها نص التعليق أدناه.', 11, False, (0x55, 0x55, 0x55))

    run(rtl_par(d.add_paragraph()), 'أولًا: النص والتوقيت', 16, True)
    tb = d.add_table(rows=1, cols=4)
    tb.style = 'Table Grid'
    tp = tb._tbl.tblPr; tp.append(OxmlElement('w:bidiVisual'))
    widths = (Cm(2.2), Cm(2.6), Cm(8.0), Cm(4.6))
    tb.autofit = False
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tp.append(lay)
    for i, w in enumerate(widths):
        tb.columns[i].width = w
    for i, h in enumerate(('التوقيت', 'المشهد', 'التعليق الصوتي', 'ما يظهر على الشاشة')):
        c = tb.rows[0].cells[i]; c.width = widths[i]
        run(rtl_par(c.paragraphs[0]), h, 11, True)
    for sc, s in zip(T['scenes'], script):
        text = ' '.join(plain(ck['text']) for ck in timing[s['id']]['clauses'])
        row = tb.add_row().cells
        vals = (f'{mm(sc["start"])}\n{mm(sc["start"] + sc["dur"])}', s['title'], text, SCREEN[s['id']])
        for i, v in enumerate(vals):
            row[i].width = widths[i]
            run(rtl_par(row[i].paragraphs[0]), v, 10.5 if i != 2 else 11.5, i == 1)

    d.add_page_break()
    run(rtl_par(d.add_paragraph()), 'ثانيًا: نص التعليق متصلًا (مشكولًا للمذيع)', 16, True)
    for s in script:
        run(rtl_par(d.add_paragraph()), s['title'], 13, True, (0x0B, 0x5C, 0x6B))
        run(rtl_par(d.add_paragraph()), s['text'], 14)

    run(rtl_par(d.add_paragraph()), 'ثالثًا: مصادر الصور', 16, True)
    for c in credits:
        run(rtl_par(d.add_paragraph()), f"• {c['caption']}: {c['title'][5:]} — {c['license']} — Wikimedia Commons", 10.5)
        if c.get('source'):
            run(d.add_paragraph(), c['source'], 9, False, (0x44, 0x44, 0x88), 'Arial')
    run(rtl_par(d.add_paragraph()), '• الرازي: رسم من مخطوطة لاتينية تعود إلى نحو 1250م (ليست صورته الحقيقية، ولا توجد له صورة).', 10.5)
    run(rtl_par(d.add_paragraph()), '• لافوازييه: نقش قديم. جابر بن حيان: لا توجد له صورة حقيقية.', 10.5)
    d.save(os.path.join(OUT, 'chemistry_script.docx'))

    md = [f'# فيلم وثائقي: الكيمياء\n', f'المدة {mm(T["total"])} • إعداد: {name} • {cls}\n', '## النص والتوقيت\n']
    for sc, s in zip(T['scenes'], script):
        md.append(f'### {mm(sc["start"])} – {mm(sc["start"] + sc["dur"])} · {s["title"]}\n')
        md.append(' '.join(plain(ck['text']) for ck in timing[s['id']]['clauses']) + '\n')
        md.append(f'*على الشاشة:* {SCREEN[s["id"]]}\n')
    open(os.path.join(OUT, 'chemistry_script.md'), 'w', encoding='utf-8').write('\n'.join(md))
    print('wrote', os.path.join(OUT, 'chemistry_script.docx'))


if __name__ == '__main__':
    main()
